import os
import sys
import argparse
from pathlib import Path
from typing import Literal

import anthropic
from pypdf import PdfReader
import requests
from bs4 import BeautifulSoup
from docx import Document
from pydantic import BaseModel


class CVOutput(BaseModel):
    cv: str
    motivatiebrief: str


def extract_pdf_text(path: str) -> str:
    reader = PdfReader(path)
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def fetch_vacancy(url: str) -> str:
    response = requests.get(url, timeout=10, headers={"User-Agent": "Mozilla/5.0"})
    response.raise_for_status()
    soup = BeautifulSoup(response.text, "html.parser")
    for tag in soup(["script", "style", "nav", "footer", "header"]):
        tag.decompose()
    lines = [l.strip() for l in soup.get_text(separator="\n").splitlines() if l.strip()]
    return "\n".join(lines)


def generate(cv_text: str, vacancy_text: str) -> CVOutput:
    api_key = os.environ.get("ANTHROPIC_API_KEY")
    if not api_key:
        raise EnvironmentError("ANTHROPIC_API_KEY is niet ingesteld.")

    client = anthropic.Anthropic(api_key=api_key)

    prompt = f"""\
Je bent een professionele loopbaancoach. Je krijgt een origineel CV en een vacaturetekst.
Jouw taak:
1. Schrijf een aangepast CV dat de meest relevante ervaring en vaardigheden voor deze vacature benadrukt. Houd de feiten exact zoals in het origineel — verzin niets. Pas alleen de volgorde, nadruk en formulering aan.
2. Schrijf een professionele, persoonlijke motivatiebrief (max. 4 alinea's) die aansluit bij de vacature en de toon van het bedrijf.

Geef beide terug in de gevraagde velden. Schrijf in dezelfde taal als de vacature.

## Origineel CV

{cv_text}

## Vacaturetekst

{vacancy_text}
"""

    response = client.messages.parse(
        model="claude-opus-4-7",
        max_tokens=4096,
        messages=[{"role": "user", "content": prompt}],
        output_format=CVOutput,
    )

    return response.parsed_output


def save_docx(output: CVOutput, output_path: str) -> None:
    doc = Document()

    doc.add_heading("Aangepast CV", level=1)
    for line in output.cv.splitlines():
        if line.strip():
            doc.add_paragraph(line)
        else:
            doc.add_paragraph("")

    doc.add_page_break()

    doc.add_heading("Motivatiebrief", level=1)
    for line in output.motivatiebrief.splitlines():
        if line.strip():
            doc.add_paragraph(line)
        else:
            doc.add_paragraph("")

    doc.save(output_path)


def main():
    parser = argparse.ArgumentParser(description="Genereer een aangepast CV en motivatiebrief.")
    parser.add_argument("cv_pdf", help="Pad naar je huidige CV (PDF)")
    parser.add_argument("vacancy_url", help="URL van de vacature")
    parser.add_argument("--output", "-o", default="output", help="Bestandsnaam voor output (zonder extensie)")
    args = parser.parse_args()

    print("CV inlezen...")
    cv_text = extract_pdf_text(args.cv_pdf)

    print("Vacature ophalen...")
    vacancy_text = fetch_vacancy(args.vacancy_url)

    print("Genereren met Claude...")
    result = generate(cv_text, vacancy_text)

    # Tekstuitvoer
    txt_path = f"{args.output}.txt"
    Path(txt_path).write_text(
        f"=== AANGEPAST CV ===\n\n{result.cv}\n\n=== MOTIVATIEBRIEF ===\n\n{result.motivatiebrief}",
        encoding="utf-8",
    )

    # Word-uitvoer
    docx_path = f"{args.output}.docx"
    save_docx(result, docx_path)

    print(f"\nKlaar!")
    print(f"  Tekst : {txt_path}")
    print(f"  Word  : {docx_path}")
    print()
    print("=== MOTIVATIEBRIEF PREVIEW ===")
    print(result.motivatiebrief[:500] + "..." if len(result.motivatiebrief) > 500 else result.motivatiebrief)


if __name__ == "__main__":
    main()
